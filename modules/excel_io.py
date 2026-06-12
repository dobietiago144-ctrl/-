"""Excel 批量导入导出"""

import io
import os
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

from config import DOC_STATUS_OPTIONS, DOC_CATEGORY_OPTIONS, REGION_OPTIONS
import database.db as db

# 导入模板表头
IMPORT_HEADERS = [
    "文件名称", "文号", "发文单位", "发布日期", "实施日期",
    "失效日期", "文件状态", "适用地区", "文件类别", "关键词",
    "业务类型", "自定义业务类型", "文件重要级别", "敏感级别", "废止依据", "替代文件", "原文链接",
    "原文件路径", "备注",
]

# 导出表头
EXPORT_HEADERS = [
    "ID", "文件名称", "文号", "发文单位", "发布日期", "实施日期",
    "失效日期", "文件状态", "适用地区", "文件类别", "关键词",
    "业务类型", "文件重要级别", "敏感级别", "原文件名", "文件大小", "文件类型",
    "摘要", "来源类型", "来源网站", "原文链接", "网页发布日期",
    "文件沿革", "是否确认", "备注", "创建时间", "更新时间",
]


def generate_import_template() -> bytes:
    """生成导入模板 Excel，供用户下载"""
    wb = Workbook()
    ws = wb.active
    ws.title = "政策文件导入模板"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for col, h in enumerate(IMPORT_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    example = [
        "XX管理办法", "自然资发〔2024〕10号", "自然资源部",
        "2024-01-15", "2024-03-01", "", "现行有效",
        "全国", "部门规章", "耕地保护,占补平衡",
        "耕地保护类", "", "一般", "公开", "", "", "", "", "",
    ]
    for col, val in enumerate(example, 1):
        ws.cell(row=2, column=col, value=val)

    widths = [30, 22, 16, 12, 12, 12, 12, 10, 12, 20, 15, 15, 12, 10, 30, 30, 30, 20, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    ws2 = wb.create_sheet("填写说明")
    notes = [
        ["字段", "说明", "可选值"],
        ["文件名称", "必填，政策文件正式名称", ""],
        ["文号", "如 自然资发〔2024〕10号", ""],
        ["发文单位", "发文机关", ""],
        ["发布日期", "格式 YYYY-MM-DD", ""],
        ["实施日期", "格式 YYYY-MM-DD", ""],
        ["失效日期", "格式 YYYY-MM-DD，不确定可留空", ""],
        ["文件状态", "从可选值中选择", "、".join(DOC_STATUS_OPTIONS)],
        ["适用地区", "全国/省/市/县", "、".join(REGION_OPTIONS)],
        ["文件类别", "从可选值中选择", "、".join(DOC_CATEGORY_OPTIONS)],
        ["关键词", "多个关键词用逗号分隔", ""],
        ["业务类型", "可多选，逗号分隔，留空自动分类", "增减挂钩类,耕地保护类,..."],
        ["自定义业务类型", "自定义业务类型，逗号分隔", "城市更新类,耕地动态平衡类"],
        ["文件重要级别", "核心/重要/一般/参考/待评估", "一般"],
        ["敏感级别", "公开/内部/敏感/涉密禁止上传", "公开"],
        ["废止依据", "说明废止该文件的新文件或依据", ""],
        ["替代文件", "替代本文件的新文件名称", ""],
        ["原文链接", "政府网站原文链接", ""],
        ["原文件路径", "主机上原文件的本地路径，可选", ""],
        ["备注", "其他补充信息", ""],
    ]
    for r, row_data in enumerate(notes, 1):
        for c, val in enumerate(row_data, 1):
            cell = ws2.cell(row=r, column=c, value=val)
            if r == 1:
                cell.font = Font(bold=True)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def parse_import_excel(file_path: str) -> dict:
    """解析导入 Excel，返回 {success: [...], errors: [...], relation_notes: [...]}"""
    try:
        df = pd.read_excel(file_path)
    except Exception as e:
        return {"success": [], "errors": [f"Excel 读取失败: {str(e)}"], "relation_notes": []}

    col_map = {}
    for i, h in enumerate(IMPORT_HEADERS):
        if h in df.columns:
            col_map[h] = h
        for col in df.columns:
            if col.strip() == h or h in col:
                col_map[h] = col
                break

    successes = []
    errors = []
    relation_notes = []

    for idx, row in df.iterrows():
        row_num = idx + 2
        try:
            title = str(row.get(col_map.get("文件名称", ""), "")).strip()
            if not title or title == "nan":
                errors.append(f"第{row_num}行：文件名称不能为空")
                continue

            doc_no = str(row.get(col_map.get("文号", ""), "")).strip()
            if doc_no and doc_no != "nan":
                existing = db.get_document_by_no(doc_no)
                if existing:
                    errors.append(f"第{row_num}行：文号「{doc_no}」已存在（《{existing['title']}》），请检查")
                    continue

            status = str(row.get(col_map.get("文件状态", ""), "")).strip()
            if status and status != "nan" and status not in DOC_STATUS_OPTIONS:
                errors.append(f"第{row_num}行：文件状态「{status}」不在预设列表中")
                continue

            abolish_basis = _safe_str(row, col_map, "废止依据")
            alt_file_name = _safe_str(row, col_map, "替代文件")

            # 检查敏感级别
            sensitivity = _safe_str(row, col_map, "敏感级别") or "公开"
            if "涉密禁止上传" in sensitivity:
                errors.append(f"第{row_num}行：敏感级别为'涉密禁止上传'，禁止导入本系统，已跳过")
                continue

            # 业务类型
            business_tags = _safe_str(row, col_map, "业务类型")
            custom_business_tags = _safe_str(row, col_map, "自定义业务类型")
            # 合并预设业务类型和自定义业务类型
            bt_set = set()
            if business_tags:
                for t in business_tags.replace("，", ",").replace("、", ",").replace(";", ",").split(","):
                    t = t.strip()
                    if t:
                        bt_set.add(t)
            if custom_business_tags:
                for t in custom_business_tags.replace("，", ",").replace("、", ",").replace(";", ",").split(","):
                    t = t.strip()
                    if t:
                        bt_set.add(t)
            if not bt_set:
                from modules.classifier import classify_business_tags
                kws = _safe_str(row, col_map, "关键词")
                auto_tags = classify_business_tags(title, kws, "")
                for t in auto_tags:
                    bt_set.add(t)
            business_tags = ",".join(sorted(bt_set))

            # 文件重要级别
            importance_level = _safe_str(row, col_map, "文件重要级别")
            if importance_level:
                from config import IMPORTANCE_LEVEL_OPTIONS
                if importance_level not in IMPORTANCE_LEVEL_OPTIONS:
                    importance_level = "待评估"  # 不在选项中时默认"待评估"
            if not importance_level:
                importance_level = "一般"

            # 原文件路径
            src_path = _safe_str(row, col_map, "原文件路径")
            file_name = ""
            file_size = 0
            file_type = ""
            file_path = ""
            if src_path and os.path.isfile(src_path):
                file_name = os.path.basename(src_path)
                file_size = os.path.getsize(src_path)
                _, ext = os.path.splitext(src_path)
                file_type = ext.lstrip(".").lower()
                from config import POLICY_FILES_DIR, SAVE_POLICY_FILE
                import shutil
                import uuid
                if SAVE_POLICY_FILE:
                    dest = os.path.join(POLICY_FILES_DIR, f"{uuid.uuid4().hex}_{file_name}")
                    shutil.copy2(src_path, dest)
                    file_path = dest

            # 构建 notes：如果替代文件未找到，将信息写入 notes
            extra_notes = _safe_str(row, col_map, "备注")

            data = {
                "title": title,
                "document_no": doc_no if doc_no != "nan" else "",
                "issuing_authority": _safe_str(row, col_map, "发文单位"),
                "publish_date": _safe_str(row, col_map, "发布日期"),
                "effective_date": _safe_str(row, col_map, "实施日期"),
                "expiry_date": _safe_str(row, col_map, "失效日期"),
                "status": status if status != "nan" else "待核实",
                "region": _safe_str(row, col_map, "适用地区"),
                "category": _safe_str(row, col_map, "文件类别"),
                "keywords": _safe_str(row, col_map, "关键词"),
                "business_tags": business_tags,
                "importance_level": importance_level,
                "sensitivity_level": sensitivity,
                "file_name": file_name,
                "file_size": file_size,
                "file_type": file_type,
                "source_type": "批量导入",
                "source_url": _safe_str(row, col_map, "原文链接"),
                "file_path": file_path,
                "notes": extra_notes,
                "confirmed": 1,
            }
            doc_id = db.create_document(data)
            successes.append({"id": doc_id, "title": title})

            # 处理替代文件和废止依据 — 自动创建 document_relations
            if alt_file_name:
                alt_file_name = alt_file_name.strip()
                # 尝试在文件库中查找替代文件
                alt_doc = db.get_document_by_title(alt_file_name)
                if not alt_doc:
                    # 尝试文号匹配
                    alt_doc = db.get_document_by_no(alt_file_name)

                if alt_doc and alt_doc["id"] != doc_id:
                    rel_type = _determine_relation_type(status)
                    rel_data = {
                        "old_document_id": doc_id,
                        "new_document_id": alt_doc["id"],
                        "relation_type": rel_type,
                        "relation_basis": abolish_basis,
                        "relation_date": datetime.now().strftime("%Y-%m-%d"),
                        "affected_scope": "全文",
                        "confidence": "待核实",
                        "notes": f"Excel批量导入自动创建：替代文件为《{alt_file_name}》",
                    }
                    db.create_relation(rel_data)
                elif alt_doc and alt_doc["id"] == doc_id:
                    relation_notes.append(f"《{title}》：替代文件指向自身，已忽略")
                else:
                    # 替代文件未找到，提示后续补建
                    note_text = f"替代文件「{alt_file_name}」未在文件库中找到，请后续补建关系"
                    if extra_notes:
                        note_text = extra_notes + "；" + note_text
                    db.update_document(doc_id, {"notes": note_text})
                    relation_notes.append(f"《{title}》：替代文件「{alt_file_name}」未在文件库中找到，请后续补建关系")

            # 如果有废止依据但没有替代文件，也写入备注
            if abolish_basis and not alt_file_name:
                note_text = f"废止依据：{abolish_basis}"
                current_doc = db.get_document(doc_id)
                if current_doc and current_doc.get("notes"):
                    note_text = current_doc["notes"] + "；" + note_text
                db.update_document(doc_id, {"notes": note_text})

        except Exception as e:
            errors.append(f"第{row_num}行：{str(e)}")

    return {"success": successes, "errors": errors, "relation_notes": relation_notes}


def _determine_relation_type(status: str) -> str:
    """根据文件状态判断关系类型"""
    status_map = {
        "已废止": "废止",
        "已失效": "废止",
        "被替代": "替代",
        "部分废止": "部分废止",
    }
    return status_map.get(status, "不确定")


def _safe_str(row, col_map: dict, key: str) -> str:
    """安全获取字符串值"""
    col = col_map.get(key, "")
    if not col or col not in row.index:
        return ""
    val = str(row[col]).strip()
    return "" if val == "nan" else val


def export_documents_to_excel(documents: list[dict]) -> bytes:
    """导出政策文件库为 Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "政策文件库"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for col, h in enumerate(EXPORT_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    field_map = {
        "ID": "id", "文件名称": "title", "文号": "document_no",
        "发文单位": "issuing_authority", "发布日期": "publish_date",
        "实施日期": "effective_date", "失效日期": "expiry_date",
        "文件状态": "status", "适用地区": "region", "文件类别": "category",
        "关键词": "keywords", "业务类型": "business_tags",
        "文件重要级别": "importance_level",
        "敏感级别": "sensitivity_level", "原文件名": "file_name",
        "文件大小": "file_size", "文件类型": "file_type",
        "摘要": "summary",
        "来源类型": "source_type", "来源网站": "source_name",
        "原文链接": "source_url", "网页发布日期": "source_publish_date",
        "文件沿革": "revision_history",
        "是否确认": "confirmed", "备注": "notes",
        "创建时间": "created_at", "更新时间": "updated_at",
    }
    for r, doc in enumerate(documents, 2):
        for c, header in enumerate(EXPORT_HEADERS, 1):
            val = doc.get(field_map.get(header, ""), "")
            if header == "是否确认":
                val = "是" if val == 1 else "否"
            ws.cell(row=r, column=c, value=val)

    for i in range(1, len(EXPORT_HEADERS) + 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 16

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def export_relations_to_excel(relations: list[dict]) -> bytes:
    """导出新旧关系为 Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "新旧关系表"

    headers = [
        "ID", "旧文件", "旧文件文号", "关系类型", "新文件", "新文件文号",
        "关系依据", "关系日期", "影响范围", "确认程度", "备注",
    ]
    fields = [
        "id", "old_title", "old_document_no", "relation_type",
        "new_title", "new_document_no", "relation_basis", "relation_date",
        "affected_scope", "confidence", "notes",
    ]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill

    for r, rel in enumerate(relations, 2):
        for c, f in enumerate(fields, 1):
            ws.cell(row=r, column=c, value=rel.get(f, ""))

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def export_review_results_to_excel(results: list[dict]) -> bytes:
    """导出审查结果为 Excel，同时展示报告引用文号和文件库正式文号"""
    wb = Workbook()
    ws = wb.active
    ws.title = "审查结果"

    headers = [
        "序号", "原文引用", "报告引用名称", "报告引用文号",
        "文件库正式名称", "文件库正式文号", "匹配方式", "匹配分数",
        "文件状态", "风险等级", "问题类型", "置信度",
        "判断依据", "建议替换文件", "修改建议",
        "出现次数", "位置", "是否核心问题", "是否需要人工确认",
    ]
    fields = [
        None, "original_reference", "recognized_title", "recognized_document_no",
        "official_title", "official_document_no", "match_method", "match_score",
        "document_status", "risk_level", "problem_type", "confidence",
        "judgment_basis", "suggested_title", "suggestion",
        "occurrence_count", "location", "is_core_issue", "need_manual_confirm",
    ]

    header_font = Font(bold=True, color="FFFFFF")
    severe_fill = PatternFill(start_color="E60000", end_color="E60000", fill_type="solid")
    high_fill = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
    mid_fill = PatternFill(start_color="FFD93D", end_color="FFD93D", fill_type="solid")
    low_fill = PatternFill(start_color="A8D8EA", end_color="A8D8EA", fill_type="solid")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill

    for r, res in enumerate(results, 2):
        for c, f in enumerate(fields, 1):
            if f is None:
                val = r - 1
            elif f == "match_score":
                val = res.get(f, 0) or 0
            elif f in ("need_manual_confirm", "is_core_issue"):
                val = "是" if res.get(f, 0) else "否"
            elif f == "occurrence_count":
                val = res.get(f, 1) or 1
            else:
                val = res.get(f, "")
            cell = ws.cell(row=r, column=c, value=val)
            # 风险等级着色
            if c == 10:  # 风险等级列
                risk = str(val)
                if risk == "严重问题":
                    cell.fill = severe_fill
                    cell.font = Font(color="FFFFFF", bold=True)
                elif risk == "高风险":
                    cell.fill = high_fill
                elif risk == "中风险":
                    cell.fill = mid_fill
                elif risk == "低风险":
                    cell.fill = low_fill

    for i in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 18

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
